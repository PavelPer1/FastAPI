import os
from fastapi import APIRouter, UploadFile, HTTPException
from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError

from auth_service.src.models.crud import register_user, get_user_by_username
from auth_service.src.schemas.user import ExcelUploadResponse

router = APIRouter()


@router.post("/register-from-excel", response_model=ExcelUploadResponse)
async def register_from_excel(file: UploadFile):
    """
    Загружает Excel файл с данными пользователей и регистрирует их.
    
    Ожидаемые колонки в Excel:
    - username (обязательно)
    - password (обязательно)
    - email (опционально)
    - full_name (опционально)
    - phone (опционально)
    - department (опционально)
    """
    if not file.filename or not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="Файл должен быть в формате Excel (.xlsx)")

    # Создаём директорию для файлов если не существует
    os.makedirs("users_files", exist_ok=True)
    file_path = f"users_files/{file.filename}"

    try:
        # Сохраняем файл
        content = await file.read()
        with open(file_path, "wb") as f:
            f.write(content)

        # Читаем Excel
        workbook = load_workbook(file_path)
        sheet = workbook.active

        # Получаем заголовки
        headers = [cell.value for cell in sheet[1] if cell.value]

        # Проверяем обязательные колонки
        required_columns = ['username', 'password']
        for col in required_columns:
            if col not in headers:
                raise HTTPException(
                    status_code=400,
                    detail=f"В Excel файле отсутствует обязательная колонка: {col}"
                )

        # Индексы колонок
        username_idx = headers.index('username')
        password_idx = headers.index('password')
        email_idx = headers.index('email') if 'email' in headers else None
        full_name_idx = headers.index('full_name') if 'full_name' in headers else None
        phone_idx = headers.index('phone') if 'phone' in headers else None
        department_idx = headers.index('department') if 'department' in headers else None

        registered_users = []
        skipped_users = []
        failed_users = []

        # Обрабатываем каждую строку
        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if not any(row):
                continue

            try:
                username = str(row[username_idx]).strip() if row[username_idx] else None
                password = str(row[password_idx]).strip() if row[password_idx] else None

                # Валидация обязательных полей
                if not username:
                    raise ValueError("Username не может быть пустым")
                if not password:
                    raise ValueError("Password не может быть пустым")
                if len(username) < 3:
                    raise ValueError("Username должен быть не менее 3 символов")
                if len(password) < 6:
                    raise ValueError("Password должен быть не менее 6 символов")

                # Проверяем существование пользователя
                existing_user = await get_user_by_username(username)
                if existing_user:
                    skipped_users.append({
                        "row": row_idx,
                        "username": username,
                        "reason": "User already exists"
                    })
                    continue

                # Получаем опциональные поля
                email = str(row[email_idx]).strip() if email_idx and row[email_idx] else None
                full_name = str(row[full_name_idx]).strip() if full_name_idx and row[full_name_idx] else None
                phone = str(row[phone_idx]).strip() if phone_idx and row[phone_idx] else None
                department = str(row[department_idx]).strip() if department_idx and row[department_idx] else None

                # Регистрируем пользователя
                await register_user(
                    username=username,
                    password=password,
                    email=email if email else None,
                    full_name=full_name,
                    phone=phone,
                    department=department
                )

                user_info = {
                    "row": row_idx,
                    "username": username,
                    "status": "success"
                }
                if email:
                    user_info["email"] = email

                registered_users.append(user_info)

            except IntegrityError as e:
                error_msg = str(e).lower()
                if "duplicate" in error_msg or "already exists" in error_msg:
                    skipped_users.append({
                        "row": row_idx,
                        "username": username if username else "Unknown",
                        "reason": "User already exists (database constraint)"
                    })
                else:
                    failed_users.append({
                        "row": row_idx,
                        "username": username if username else "Unknown",
                        "error": f"Database integrity error: {str(e)[:200]}"
                    })
            except Exception as e:
                failed_users.append({
                    "row": row_idx,
                    "username": username if username else "Unknown",
                    "error": str(e)
                })

        # Удаляем временный файл
        if os.path.exists(file_path):
            os.remove(file_path)

        # Формируем ответ
        total_processed = len(registered_users) + len(skipped_users) + len(failed_users)
        has_errors = len(failed_users) > 0

        return ExcelUploadResponse(
            success=not has_errors,
            message=f"Новых: {len(registered_users)}, Пропущено: {len(skipped_users)}, Ошибок: {len(failed_users)}",
            registered_users=registered_users,
            skipped_users=skipped_users,
            failed_users=failed_users,
            statistics={
                "total_processed": total_processed,
                "new_registrations": len(registered_users),
                "already_existed": len(skipped_users),
                "errors": len(failed_users),
                "success_rate": f"{(len(registered_users) / total_processed * 100):.1f}%" if total_processed > 0 else "0%"
            }
        )

    except HTTPException:
        if os.path.exists(file_path):
            os.remove(file_path)
        raise
    except Exception as e:
        if os.path.exists(file_path):
            os.remove(file_path)
        raise HTTPException(
            status_code=500,
            detail=f"Ошибка при обработке файла: {str(e)}"
        )
